"""Write data/import-sample.xlsx, the file the Import page offers as its sample.

    .venv/bin/python -m tests.tools.make_import_sample

Launch dates are real Excel dates counted from the day you run this, so run it again before a demo
if the file is more than a few days old (old dates show "launch has passed" warnings, and a year
later they become errors). The sheet is built to show every outcome on a freshly reset demo as
Maya Chen: clean rows, rows that trip flags (R1 to R7), rows that need fixing, and one exact repeat
of a seed item that shows as "Already exists".
"""
import json
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "import-sample.xlsx"
HEADERS = ("Title", "Product", "Channel", "Launch date", "Copy", "Notes")


def seed_item(seed_id):
    items = json.loads((ROOT / "data" / "seed.json").read_text())["submissions"]
    return next(i for i in items if i["seedId"] == seed_id)


def rows(today):
    def when(days):
        return today + timedelta(days=days)

    repeat = seed_item(2)  # Maya's "Cashback card paid social", word for word
    return [
        # Ready, no flags.
        ("Auto loan referral email", "Loan", "Email", when(21),
         "Subject: Thinking about a new car?\n\nHi there,\n\nA ClearPath auto loan can help you drive away "
         "sooner. Check your options in a few minutes, and choose the term that fits your budget. All loans "
         "are subject to credit approval.\n\nThe ClearPath Team",
         "Monthly referral send. Same layout as last quarter."),
        ("Card balance transfer reminder", "Card", "Email", when(14),
         "Subject: Move a balance and simplify your payments\n\nHi there,\n\nA balance transfer can put your "
         "payments in one place. Intro offer, then 19.99% APR. Offer is subject to credit approval.\n\n"
         "The ClearPath Team", None),
        # Ready, with flags: R1, R2 and R5.
        ("Spring loan promo email", "Loan", "Email", when(7),
         "Subject: Spring cash, fast\n\nGuaranteed approval on every personal loan, with rates as low as "
         "5.99%. Funds can reach your account the next day.\n\nApply today.",
         "Marketing asked for the 'guaranteed' line to stay."),
        # Ready, with flags: R5, R6 and R7 (short format).
        ("Travel card paid social ad", "Card", "Paid social", when(4),
         "Earn 2x points on every trip. Act now and start earning before the holidays. #ClearPath",
         "Single image, two lines of copy."),
        # Ready, with flags: R3, R4 and R7.
        ("Mortgage display banner", "Mortgage", "Display", when(10),
         "You're approved to start house hunting. Find your rate today.", None),
        # Needs fixing: no title, unknown product, launch date typed as words, copy far too long.
        (None, "Loan", "Email", when(9),
         "Subject: A note about your account\n\nThis is a long enough message to be real copy.", None),
        ("Crypto-backed loan teaser", "Crypto", "Email", when(9),
         "Borrow against what you hold. Subject to credit approval.", None),
        ("Holiday card offer", "Card", "Email", "next Friday",
         "Give yourself a break this season. APR 21.99%. Subject to credit approval.", None),
        ("Everything about the loan (draft)", "Loan", "Email", when(12), "Lorem ipsum dolor sit amet. " * 400,
         "Pasted the whole brief by mistake."),
        # Already exists: the same item the demo starts with.
        (repeat["title"], "Card", "Paid social", when(2), repeat["versions"][0]["copy"], None),
    ]


def build(today=None):
    today = today or date.today()
    wb = Workbook()
    ws = wb.active
    ws.title = "Submissions"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows(today):
        ws.append(row)
    for cell in ws["D"][1:]:
        if isinstance(cell.value, date):
            cell.number_format = "yyyy-mm-dd"
    for index, width in enumerate((34, 12, 14, 13, 70, 36), 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    return wb


if __name__ == "__main__":
    build().save(OUT)
    print("wrote", OUT.relative_to(ROOT))
