import io
import zipfile
from datetime import date, datetime, timezone

import pytest
from openpyxl import Workbook

from app import db, submit, xlsx_import
from app.xlsx_import import ImportFileError

TODAY = date(2026, 10, 7)
NOW = datetime(2026, 10, 7, 9, 0, tzinfo=timezone.utc)
HEAD = ["Title", "Product", "Channel", "Launch date", "Copy", "Notes"]
COPY = "Apply in minutes. Subject to credit approval."


def good(**over):
    row = {"Title": "Spring email", "Product": "Loan", "Channel": "Email", "Launch date": date(2026, 11, 2),
           "Copy": COPY, "Notes": "n"}
    row.update(over)
    return [row[h] for h in HEAD]


def book(rows, head=HEAD, sheets=None):
    wb = Workbook()
    ws = wb.active
    if head is not None:
        ws.append(head)
    for r in rows:
        ws.append(r)
    for name in sheets or ():
        wb.create_sheet(name).append(["Title"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def rows_of(data, **kw):
    return xlsx_import.parse_workbook(data, **kw)


def fields_of(data):
    return xlsx_import.normalize_row(rows_of(data)[0]["fields"])


def code(data, **kw):
    with pytest.raises(ImportFileError) as exc:
        rows_of(data, **kw)
    return exc.value.code


# ---- the file as a whole ----

def test_reads_rows_and_numbers_them_by_sheet_row():
    rows = rows_of(book([good(), good(Title="Second")]))
    assert [r["row"] for r in rows] == [2, 3] and rows[1]["fields"]["title"] == "Second"


@pytest.mark.parametrize("data, expected", [
    (b"", "empty"),
    (b"Title,Product\nx,y\n", "not_xlsx"),          # a .csv renamed to .xlsx
    (b"\xd0\xcf\x11\xe0" + b"0" * 600, "not_xlsx"),  # an old .xls or an encrypted workbook (OLE container)
    (b"PK\x03\x04 not really a zip", "not_xlsx"),
])
def test_files_that_are_not_xlsx(data, expected):
    assert code(data) == expected


def test_macro_workbook_is_refused():
    plain = book([good()])
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(plain)) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            dst.writestr(info, src.read(info))
        dst.writestr("xl/vbaProject.bin", b"x")
    assert code(out.getvalue()) == "macros"


def test_zip_bomb_is_refused_before_it_is_opened():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<x/>")
        z.writestr("xl/workbook.xml", "<x/>")
        z.writestr("xl/worksheets/sheet1.xml", b"0" * (xlsx_import.MAX_UNZIPPED_BYTES + 1))
    assert len(out.getvalue()) < 100_000
    assert code(out.getvalue()) == "too_big"


def test_zip_without_workbook_parts_is_not_xlsx():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("hello.txt", "hi")
    assert code(out.getvalue()) == "not_xlsx"


def test_headers_only_and_empty_sheet_say_no_rows():
    assert code(book([])) == "no_rows"
    assert code(book([], head=None)) == "no_rows"


def test_blank_rows_are_skipped_and_not_counted():
    rows = rows_of(book([[None] * 6, good(), ["", " ", None], good(Title="Two"), [None] * 6]))
    assert [r["row"] for r in rows] == [3, 5]


def test_blank_lines_above_the_header_are_skipped():
    wb = Workbook()
    ws = wb.active
    ws.append([None, None])
    ws.append(HEAD)
    ws.append(good())
    out = io.BytesIO()
    wb.save(out)
    assert rows_of(out.getvalue())[0]["row"] == 3


def test_missing_columns_are_named_from_our_fixed_list():
    data = book([["a", "b"]], head=["Title", "Something else"])
    with pytest.raises(ImportFileError) as exc:
        rows_of(data)
    assert exc.value.columns == ("product", "channel", "launch_date", "copy")
    assert "Product, Channel, Launch date, Copy" in exc.value.message
    assert "Something else" not in exc.value.message


def test_channel_column_is_optional_when_the_channel_is_forced():
    head = ["Title", "Product", "Launch date", "Copy"]
    rows = rows_of(book([["T", "Loan", date(2026, 11, 2), COPY]], head=head), need_channel=False)
    assert len(rows) == 1


def test_headers_match_by_case_spacing_and_alias_and_order_is_free():
    head = ["  COPY ", "launch_date", "Channel", "product", "Title", "Notes for reviewer", "Extra"]
    data = book([[COPY, date(2026, 11, 2), "Email", "Loan", "T", "n", "ignored"]], head=head)
    f = fields_of(data)
    assert f["title"] == "T" and f["notes"] == "n" and f["copy"] == COPY and f["launch_date"] == "2026-11-02"


def test_first_of_two_duplicate_headers_wins():
    data = book([["First", "Second", "Loan", "Email", date(2026, 11, 2), COPY]],
                head=["Title", "Title", "Product", "Channel", "Launch date", "Copy"])
    assert rows_of(data)[0]["fields"]["title"] == "First"


def test_only_the_first_sheet_is_read():
    rows = rows_of(book([good()], sheets=["Other"]))
    assert len(rows) == 1


def test_row_cap_is_enforced():
    ok = book([good(Title="T%d" % i) for i in range(xlsx_import.MAX_ROWS)])
    assert len(rows_of(ok)) == xlsx_import.MAX_ROWS
    over = book([good(Title="T%d" % i) for i in range(xlsx_import.MAX_ROWS + 1)])
    assert code(over) == "too_many"


def test_a_sheet_of_blank_rows_is_cut_off():
    wb = Workbook()
    ws = wb.active
    ws.append(HEAD)
    ws.cell(row=xlsx_import.MAX_SCANNED_ROWS + 5, column=1, value="x")
    out = io.BytesIO()
    wb.save(out)
    assert code(out.getvalue()) == "too_many"


def test_formula_without_a_saved_value_reads_as_blank():
    data = book([good(Title="=1+1")])  # openpyxl saves no cached result
    assert rows_of(data)[0]["fields"]["title"] is None
    assert xlsx_import.normalize_row(rows_of(data)[0]["fields"])["title"] is None


def test_messages_are_fixed_text():
    for key in xlsx_import.FILE_MESSAGES:
        assert ImportFileError(key).message == xlsx_import.FILE_MESSAGES[key]


# ---- cells ----

@pytest.mark.parametrize("cell, expected", [
    (date(2026, 11, 2), "2026-11-02"),
    (datetime(2026, 11, 2, 15, 30), "2026-11-02"),
    ("2026-11-02", "2026-11-02"),
    ("  2026-11-02 ", "2026-11-02"),
    ("11/02/2026", "11/02/2026"),   # left for the validator to refuse
    (46328, None),                   # an Excel serial number: refused, with the YYYY-MM-DD message
    (None, None),
])
def test_launch_date_cells(cell, expected):
    assert xlsx_import.normalize_row({"launch_date": cell})["launch_date"] == expected


@pytest.mark.parametrize("cell, product", [
    ("Loan", "loan"), ("loan", "loan"), ("  CARD ", "card"), ("Mortgage", "mortgage"), ("Crypto", "Crypto"),
    ("", ""), (None, None), (7, 7)])
def test_product_labels_and_values(cell, product):
    assert xlsx_import.normalize_row({"product": cell})["product"] == product


@pytest.mark.parametrize("cell, channel", [
    ("Paid social", "paid_social"), ("paid_social", "paid_social"), ("PAID  SOCIAL", "paid_social"),
    ("Affiliate page", "affiliate_page"), ("Email", "email"), ("Display", "display"), ("Radio", "Radio")])
def test_channel_labels_and_values(cell, channel):
    assert xlsx_import.normalize_row({"channel": cell})["channel"] == channel


def test_forced_channel_beats_the_sheet():
    assert xlsx_import.normalize_row({"channel": "Email"}, force_channel="affiliate_page")["channel"] == "affiliate_page"


def test_numbers_in_text_columns_become_text():
    f = xlsx_import.normalize_row({"title": 2026, "copy": 1.0, "notes": 2.5})
    assert (f["title"], f["copy"], f["notes"]) == ("2026", "1", "2.5")
    assert xlsx_import.normalize_row({"title": True})["title"] is True  # refused by the validator


def validate(**cells):
    f = xlsx_import.normalize_row(cells)
    return submit.validate_submission(f, TODAY)


@pytest.mark.parametrize("cells, field", [
    ({"title": "", "product": "Loan", "channel": "Email", "launch_date": date(2026, 11, 2), "copy": COPY}, "title"),
    ({"title": "​ ​", "product": "Loan", "channel": "Email", "launch_date": date(2026, 11, 2), "copy": COPY}, "title"),
    ({"title": "T" * 121, "product": "Loan", "channel": "Email", "launch_date": date(2026, 11, 2), "copy": COPY}, "title"),
    ({"title": "T", "product": "Loan", "channel": "Email", "launch_date": date(2029, 1, 1), "copy": COPY}, "launch_date"),
    ({"title": "T", "product": "Loan", "channel": "Email", "launch_date": date(2024, 1, 1), "copy": COPY}, "launch_date"),
    ({"title": "T", "product": "Loan", "channel": "Email", "launch_date": date(2026, 11, 2), "copy": "x" * 10_001}, "copy"),
    ({"title": "T", "product": "Loan", "channel": "Email", "launch_date": date(2026, 11, 2), "copy": COPY,
      "notes": "n" * 2_001}, "notes"),
    ({"title": "T", "product": "Loan", "channel": "Email", "launch_date": date(2026, 11, 2), "copy": "a\x00b"}, "copy"),
])
def test_bad_cells_give_the_same_messages_as_the_form(cells, field):
    _, errors, _ = validate(**cells)
    assert list(errors) == [field] and errors[field] in submit.MESSAGES.values()


def test_past_launch_is_a_warning_not_an_error():
    clean, errors, warnings = validate(title="T", product="Loan", channel="Email", launch_date=date(2026, 9, 1),
                                       copy=COPY)
    assert not errors and warnings == ["launch_past"]


def test_odd_text_is_kept_as_text():
    clean, errors, _ = validate(title="=SUM(A1)", product="Loan", channel="Email", launch_date=date(2026, 11, 2),
                                copy="@home +1 - 2\nline two \U0001F600 {{ 7*7 }} <b>x</b>")
    assert not errors and clean["title"] == "=SUM(A1)" and "{{ 7*7 }}" in clean["copy"]


# ---- preview and import against the database ----

@pytest.fixture
def seeded(conn):
    return conn


def sheet(*rows):
    return rows_of(book(list(rows)))


def test_preview_statuses_flags_and_no_writes(conn):
    before = conn.execute("SELECT count(*) FROM submission").fetchone()[0]
    rows = sheet(good(), good(Title="Same"), good(Title="Same"), good(Title=None),
                 good(Title="Flagged", Copy="Guaranteed approval for all."))
    results = xlsx_import.preview_rows(conn, rows, TODAY, "Maya Chen")
    assert [r["status"] for r in results] == ["ok", "ok", "duplicate", "error", "ok"]
    assert results[2]["existing_id"] is None  # a repeat inside the file, not in the database
    assert "R1" in [c["rule_id"] for c in results[4]["cards"]]
    assert conn.execute("SELECT count(*) FROM submission").fetchone()[0] == before == 0


def test_preview_and_import_agree_about_existing_items(conn):
    rows = sheet(good())
    first = xlsx_import.import_rows(conn, rows, "Maya Chen", NOW)
    assert first == {"created": 1, "skipped": 0, "failed": 0, "stopped": False}
    assert xlsx_import.preview_rows(conn, rows, TODAY, "Maya Chen")[0]["status"] == "duplicate"
    assert xlsx_import.preview_rows(conn, rows, TODAY, "Jordan Lee")[0]["status"] == "ok"  # theirs, not Maya's
    assert xlsx_import.import_rows(conn, rows, "Maya Chen", NOW)["skipped"] == 1


def test_import_creates_normal_submissions(conn):
    xlsx_import.import_rows(conn, sheet(good(Copy="Guaranteed approval, rates as low as 5%.")), "Maya Chen", NOW)
    s = conn.execute("SELECT * FROM submission").fetchone()
    assert (s["status"], s["current_version"], s["submitted_by"], s["launch_date"]) == (
        "new", 1, "Maya Chen", "2026-11-02")
    assert s["created_at"] == "2026-10-07T09:00:00Z"
    assert conn.execute("SELECT count(*) FROM flag").fetchone()[0] >= 2


def test_submitted_by_column_is_ignored(conn):
    data = book([good() + ["Jordan Lee"]], head=HEAD + ["Submitted by"])
    xlsx_import.import_rows(conn, rows_of(data), "Sam Patel", NOW)
    assert conn.execute("SELECT submitted_by FROM submission").fetchone()[0] == "Sam Patel"


def test_import_keeps_earlier_rows_when_a_later_row_fails(conn):
    rows = sheet(good(Title="A"), good(Title=None), good(Title="B"))
    summary = xlsx_import.import_rows(conn, rows, "Maya Chen", NOW)
    assert (summary["created"], summary["failed"]) == (2, 1)
    assert conn.execute("SELECT count(*) FROM submission").fetchone()[0] == 2


def test_import_stops_at_capacity(conn, monkeypatch):
    monkeypatch.setattr(submit, "MAX_SUBMISSIONS", 2)
    rows = sheet(*[good(Title="T%d" % i) for i in range(4)])
    summary = xlsx_import.import_rows(conn, rows, "Maya Chen", NOW)
    assert summary == {"created": 2, "skipped": 0, "failed": 0, "stopped": True}
    assert conn.execute("SELECT count(*) FROM submission").fetchone()[0] == 2


def test_import_with_an_unknown_submitter_fails_every_row(conn):
    summary = xlsx_import.import_rows(conn, sheet(good()), "Mallory", NOW)
    assert summary["failed"] == 1 and summary["created"] == 0


# ---- the sample file in the repo ----

def test_sample_file_exists_and_previews_as_designed(client):
    rows = xlsx_import.parse_workbook(xlsx_import.read_sample())
    with db.connect() as c:
        results = xlsx_import.preview_rows(c, rows, TODAY, "Maya Chen")
    assert [r["status"] for r in results] == ["ok"] * 5 + ["error"] * 4 + ["duplicate"]
    fired = {c["rule_id"] for r in results for c in r["cards"]}
    assert fired == {"R1", "R2", "R3", "R4", "R5", "R6", "R7"}
    assert results[0]["cards"] == [] and results[1]["cards"] == []  # clean and near-miss rows
    assert results[9]["existing_id"] == 2
